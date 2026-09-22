import os

from pypdf import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt"
}


def extract_pdf(file_path):

    text = []

    reader = PdfReader(file_path)

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text.append(page_text)

    return "\n".join(text)


def extract_docx(file_path):

    document = Document(file_path)

    text = []

    # Normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text.append(
                paragraph.text.strip()
            )

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:

                text.append(
                    " | ".join(row_text)
                )

    return "\n".join(text)


def extract_txt(file_path):

    with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
    ) as file:

        return file.read()


def clean_text(text):

    if not text:
        return ""

    lines = []

    for line in text.splitlines():

        line = line.strip()

        if line:

            lines.append(line)

    return "\n".join(lines)


def extract_text(file_path):

    if not os.path.exists(file_path):

        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension not in SUPPORTED_EXTENSIONS:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":

        text = extract_pdf(
            file_path
        )

    elif extension == ".docx":

        text = extract_docx(
            file_path
        )

    elif extension == ".txt":

        text = extract_txt(
            file_path
        )

    else:

        text = ""

    return clean_text(
        text
    )


if __name__ == "__main__":

    file_path = input(
        "Enter document path: "
    ).strip()

    try:

        text = extract_text(
            file_path
        )

        print("\n" + "=" * 70)
        print("📄 EXTRACTED DOCUMENT TEXT")
        print("=" * 70)

        print(text)

        print("\n" + "=" * 70)

        print(
            "Characters extracted:",
            len(text)
        )

    except Exception as error:

        print(
            "\n❌ Error:",
            error
        )